import openpyxl
import sqlite3
import re
import time
import os
import sys
from config import DB_PATH, EXCEL_SOURCE_PATH
from database import init_db, get_db_connection

def extract_id_name(text):
    """Extract standard Telcel Site ID (e.g., CA0249, JL4089) from text."""
    if not text:
        return ""
    s = str(text).strip()
    m = re.search(r'([A-Z]{2}\d{4})', s.upper())
    if m:
        return m.group(1)
    return s[:6] if len(s) >= 6 else s

def str_val(val):
    if val is None:
        return ""
    return str(val).strip()

def import_gsm(wb, conn, batch_size=5000):
    sheet_name = "GSM_CELL_REPORT"
    if sheet_name not in wb.sheetnames:
        print(f"Sheet {sheet_name} not found.")
        return 0
    sheet = wb[sheet_name]
    cursor = conn.cursor()
    cursor.execute("DELETE FROM gsm_cells;")
    
    rows_to_insert = []
    total = 0
    t0 = time.time()
    print(f"Importing {sheet_name}...")
    
    for i, row in enumerate(sheet.iter_rows(values_only=True)):
        if i == 0 or not any(row):  # skip header or empty
            continue
        ne_name = str_val(row[0])
        site_name = str_val(row[2]) if len(row) > 2 else ""
        id_name = extract_id_name(site_name or ne_name)
        
        # pad row to 16 cols
        r = list(row) + [None] * (16 - len(row))
        record = (
            id_name,
            str_val(r[0]),  # ne_name
            str_val(r[1]),  # site_index
            str_val(r[2]),  # site_name
            str_val(r[3]),  # cell_index
            str_val(r[4]),  # cell_name
            str_val(r[5]),  # activity_status
            str_val(r[6]),  # ci
            str_val(r[7]),  # basei
            str_val(r[8]),  # ni
            str_val(r[9]),  # bcchno
            str_val(r[10]), # freq_seg
            str_val(r[11]), # blk_status
            str_val(r[12]), # hop_hsn
            str_val(r[13]), # hop_tsc
            str_val(r[14]), # hop_index
            str_val(r[15])  # lac
        )
        rows_to_insert.append(record)
        if len(rows_to_insert) >= batch_size:
            cursor.executemany("""
                INSERT INTO gsm_cells (
                    id_name, ne_name, site_index, site_name, cell_index, cell_name,
                    activity_status, ci, basei, ni, bcchno, freq_seg, blk_status,
                    hop_hsn, hop_tsc, hop_index, lac
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?);
            """, rows_to_insert)
            conn.commit()
            total += len(rows_to_insert)
            rows_to_insert.clear()

    if rows_to_insert:
        cursor.executemany("""
            INSERT INTO gsm_cells (
                id_name, ne_name, site_index, site_name, cell_index, cell_name,
                activity_status, ci, basei, ni, bcchno, freq_seg, blk_status,
                hop_hsn, hop_tsc, hop_index, lac
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?);
        """, rows_to_insert)
        conn.commit()
        total += len(rows_to_insert)

    print(f"Imported {total} rows from {sheet_name} in {time.time()-t0:.2f}s")
    return total

def import_umts(wb, conn, batch_size=5000):
    sheet_name = "UMTS_CELL_REPORT"
    if sheet_name not in wb.sheetnames:
        print(f"Sheet {sheet_name} not found.")
        return 0
    sheet = wb[sheet_name]
    cursor = conn.cursor()
    cursor.execute("DELETE FROM umts_cells;")
    
    rows_to_insert = []
    total = 0
    t0 = time.time()
    print(f"Importing {sheet_name}...")
    
    for i, row in enumerate(sheet.iter_rows(values_only=True)):
        if i == 0 or not any(row):
            continue
        nodeb_name = str_val(row[2]) if len(row) > 2 else ""
        ne_name = str_val(row[0])
        id_name = extract_id_name(nodeb_name or ne_name)
        
        r = list(row) + [None] * (18 - len(row))
        record = (
            id_name,
            str_val(r[0]),  # ne_name
            str_val(r[1]),  # nodeb_id
            str_val(r[2]),  # nodeb_name
            str_val(r[3]),  # cell_id
            str_val(r[4]),  # cell_name
            str_val(r[5]),  # rnc_connection_status
            str_val(r[6]),  # activity_status
            str_val(r[7]),  # blk_status
            str_val(r[8]),  # lac
            str_val(r[9]),  # sac
            str_val(r[10]), # rac
            str_val(r[11]), # ul_freq
            str_val(r[12]), # dl_freq
            str_val(r[13]), # max_power
            str_val(r[14]), # cell_cbs_state
            str_val(r[15]), # cell_mbms_state
            str_val(r[16]), # hsdpa_opstate
            str_val(r[17])  # hsupa_opstate
        )
        rows_to_insert.append(record)
        if len(rows_to_insert) >= batch_size:
            cursor.executemany("""
                INSERT INTO umts_cells (
                    id_name, ne_name, nodeb_id, nodeb_name, cell_id, cell_name,
                    rnc_connection_status, activity_status, blk_status, lac, sac, rac,
                    ul_freq, dl_freq, max_power, cell_cbs_state, cell_mbms_state,
                    hsdpa_opstate, hsupa_opstate
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?);
            """, rows_to_insert)
            conn.commit()
            total += len(rows_to_insert)
            rows_to_insert.clear()

    if rows_to_insert:
        cursor.executemany("""
            INSERT INTO umts_cells (
                id_name, ne_name, nodeb_id, nodeb_name, cell_id, cell_name,
                rnc_connection_status, activity_status, blk_status, lac, sac, rac,
                ul_freq, dl_freq, max_power, cell_cbs_state, cell_mbms_state,
                hsdpa_opstate, hsupa_opstate
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?);
        """, rows_to_insert)
        conn.commit()
        total += len(rows_to_insert)

    print(f"Imported {total} rows from {sheet_name} in {time.time()-t0:.2f}s")
    return total

def import_lte(wb, conn, batch_size=5000):
    sheet_name = "LTE_CELL_REPORT"
    if sheet_name not in wb.sheetnames:
        print(f"Sheet {sheet_name} not found.")
        return 0
    sheet = wb[sheet_name]
    cursor = conn.cursor()
    cursor.execute("DELETE FROM lte_cells;")
    
    rows_to_insert = []
    total = 0
    t0 = time.time()
    print(f"Importing {sheet_name}...")
    
    for i, row in enumerate(sheet.iter_rows(values_only=True)):
        if i == 0 or not any(row):
            continue
        lte_ne = str_val(row[4]) if len(row) > 4 else ""
        func_name = str_val(row[5]) if len(row) > 5 else ""
        id_name = extract_id_name(lte_ne or func_name)
        
        r = list(row) + [None] * (16 - len(row))
        record = (
            id_name,
            str_val(r[0]),  # subarea
            str_val(r[1]),  # rat
            str_val(r[2]),  # operator
            str_val(r[3]),  # enodeb_id
            str_val(r[4]),  # lte_ne_name
            str_val(r[5]),  # enodeb_function_name
            str_val(r[6]),  # ne_connection_status
            str_val(r[7]),  # cell_id
            str_val(r[8]),  # cell_name
            str_val(r[9]),  # local_cell_id
            str_val(r[10]), # tac
            str_val(r[11]), # frequency_band
            str_val(r[12]), # administrative_status
            str_val(r[13]), # activation_status
            str_val(r[14]), # operating_status
            str_val(r[15])  # availability_status
        )
        rows_to_insert.append(record)
        if len(rows_to_insert) >= batch_size:
            cursor.executemany("""
                INSERT INTO lte_cells (
                    id_name, subarea, rat, operator, enodeb_id, lte_ne_name,
                    enodeb_function_name, ne_connection_status, cell_id, cell_name,
                    local_cell_id, tac, frequency_band, administrative_status,
                    activation_status, operating_status, availability_status
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?);
            """, rows_to_insert)
            conn.commit()
            total += len(rows_to_insert)
            rows_to_insert.clear()

    if rows_to_insert:
        cursor.executemany("""
            INSERT INTO lte_cells (
                id_name, subarea, rat, operator, enodeb_id, lte_ne_name,
                enodeb_function_name, ne_connection_status, cell_id, cell_name,
                local_cell_id, tac, frequency_band, administrative_status,
                activation_status, operating_status, availability_status
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?);
        """, rows_to_insert)
        conn.commit()
        total += len(rows_to_insert)

    print(f"Imported {total} rows from {sheet_name} in {time.time()-t0:.2f}s")
    return total

def import_nr(wb, conn, batch_size=5000):
    sheet_name = "NR_CELL_REPORT"
    if sheet_name not in wb.sheetnames:
        print(f"Sheet {sheet_name} not found.")
        return 0
    sheet = wb[sheet_name]
    cursor = conn.cursor()
    cursor.execute("DELETE FROM nr_cells;")
    
    rows_to_insert = []
    total = 0
    t0 = time.time()
    print(f"Importing {sheet_name}...")
    
    for i, row in enumerate(sheet.iter_rows(values_only=True)):
        if i == 0 or not any(row):
            continue
        nr_ne = str_val(row[4]) if len(row) > 4 else ""
        func_name = str_val(row[5]) if len(row) > 5 else ""
        id_name = extract_id_name(nr_ne or func_name)
        
        r = list(row) + [None] * (16 - len(row))
        record = (
            id_name,
            str_val(r[0]),  # subarea
            str_val(r[1]),  # rat
            str_val(r[2]),  # operator
            str_val(r[3]),  # gnodeb_id
            str_val(r[4]),  # nr_ne_name
            str_val(r[5]),  # gnodeb_function_name
            str_val(r[6]),  # ne_connection_status
            str_val(r[7]),  # nr_cell_id
            str_val(r[8]),  # cell_name
            str_val(r[9]),  # cell_id
            str_val(r[10]), # tac
            str_val(r[11]), # frequency_band
            str_val(r[12]), # administrative_status
            str_val(r[13]), # activation_status
            str_val(r[14]), # operating_status
            str_val(r[15])  # availability_status
        )
        rows_to_insert.append(record)
        if len(rows_to_insert) >= batch_size:
            cursor.executemany("""
                INSERT INTO nr_cells (
                    id_name, subarea, rat, operator, gnodeb_id, nr_ne_name,
                    gnodeb_function_name, ne_connection_status, nr_cell_id, cell_name,
                    cell_id, tac, frequency_band, administrative_status,
                    activation_status, operating_status, availability_status
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?);
            """, rows_to_insert)
            conn.commit()
            total += len(rows_to_insert)
            rows_to_insert.clear()

    if rows_to_insert:
        cursor.executemany("""
            INSERT INTO nr_cells (
                id_name, subarea, rat, operator, gnodeb_id, nr_ne_name,
                gnodeb_function_name, ne_connection_status, nr_cell_id, cell_name,
                cell_id, tac, frequency_band, administrative_status,
                activation_status, operating_status, availability_status
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?);
        """, rows_to_insert)
        conn.commit()
        total += len(rows_to_insert)

    print(f"Imported {total} rows from {sheet_name} in {time.time()-t0:.2f}s")
    return total

def import_rru_catalog(wb, conn):
    sheet_name = "RRU Data Base"
    if sheet_name not in wb.sheetnames:
        return 0
    sheet = wb[sheet_name]
    cursor = conn.cursor()
    cursor.execute("DELETE FROM rru_catalog;")
    
    rows = []
    for i, row in enumerate(sheet.iter_rows(values_only=True)):
        if i == 0 or not any(row):
            continue
        r = list(row) + [None] * (4 - len(row))
        rows.append((str_val(r[0]), str_val(r[1]), str_val(r[2]), str_val(r[3])))
    
    cursor.executemany("""
        INSERT INTO rru_catalog (radio, bandas_soportadas, potencia, conector)
        VALUES (?, ?, ?, ?);
    """, rows)
    conn.commit()
    print(f"Imported {len(rows)} rows into rru_catalog")
    return len(rows)

def import_rru_items(wb, conn, batch_size=5000):
    sheet_name = "RRU Filtro"
    if sheet_name not in wb.sheetnames:
        sheet_name = "NE items"
    if sheet_name not in wb.sheetnames:
        return 0
    sheet = wb[sheet_name]
    cursor = conn.cursor()
    cursor.execute("DELETE FROM rru_items;")
    
    rows_to_insert = []
    total = 0
    t0 = time.time()
    print(f"Importing {sheet_name}...")
    
    for i, row in enumerate(sheet.iter_rows(values_only=True)):
        if i == 0 or not any(row):
            continue
        ne_name = str_val(row[0]) if len(row) > 0 else ""
        id_name = str_val(row[1]) if len(row) > 1 and row[1] else extract_id_name(ne_name)
        
        r = list(row) + [None] * (12 - len(row))
        record = (
            ne_name,
            id_name,
            str_val(r[2]),  # board_name
            str_val(r[3]),  # manufacturer_data
            str_val(r[4]),  # modelo_rru
            str_val(r[5]),  # bandas_soportadas
            str_val(r[6]),  # blank_col
            str_val(r[7]),  # potencia
            str_val(r[8]),  # no_serie
            str_val(r[9]),  # conector
            str_val(r[10]), # subrack
            str_val(r[11])  # aux
        )
        rows_to_insert.append(record)
        if len(rows_to_insert) >= batch_size:
            cursor.executemany("""
                INSERT INTO rru_items (
                    ne_name, id_name, board_name, manufacturer_data, modelo_rru,
                    bandas_soportadas, blank_col, potencia, no_serie, conector, subrack, aux
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?);
            """, rows_to_insert)
            conn.commit()
            total += len(rows_to_insert)
            rows_to_insert.clear()

    if rows_to_insert:
        cursor.executemany("""
            INSERT INTO rru_items (
                ne_name, id_name, board_name, manufacturer_data, modelo_rru,
                bandas_soportadas, blank_col, potencia, no_serie, conector, subrack, aux
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?);
        """, rows_to_insert)
        conn.commit()
        total += len(rows_to_insert)

    print(f"Imported {total} rows from {sheet_name} in {time.time()-t0:.2f}s")
    return total

def import_inventory(wb, conn, batch_size=10000):
    sheet_name = "Inventory Board"
    if sheet_name not in wb.sheetnames:
        return 0
    sheet = wb[sheet_name]
    cursor = conn.cursor()
    cursor.execute("DELETE FROM inventory_boards;")
    
    rows_to_insert = []
    total = 0
    t0 = time.time()
    print(f"Importing {sheet_name} (this is large, streaming chunks)...")
    
    for i, row in enumerate(sheet.iter_rows(values_only=True)):
        if i == 0 or not any(row):
            continue
        ne_name = str_val(row[2]) if len(row) > 2 else ""
        id_name = extract_id_name(ne_name)
        
        r = list(row) + [None] * (41 - len(row))
        record = (
            id_name,
            str_val(r[0]), str_val(r[1]), str_val(r[2]), str_val(r[3]), str_val(r[4]),
            str_val(r[5]), str_val(r[6]), str_val(r[7]), str_val(r[8]), str_val(r[9]),
            str_val(r[10]), str_val(r[11]), str_val(r[12]), str_val(r[13]), str_val(r[14]),
            str_val(r[15]), str_val(r[16]), str_val(r[17]), str_val(r[18]), str_val(r[19]),
            str_val(r[20]), str_val(r[21]), str_val(r[22]), str_val(r[23]), str_val(r[24]),
            str_val(r[25]), str_val(r[26]), str_val(r[27]), str_val(r[28]), str_val(r[29]),
            str_val(r[30]), str_val(r[31]), str_val(r[32]), str_val(r[33]), str_val(r[34]),
            str_val(r[35]), str_val(r[36]), str_val(r[37]), str_val(r[38]), str_val(r[39]),
            str_val(r[40])
        )
        rows_to_insert.append(record)
        if len(rows_to_insert) >= batch_size:
            cursor.executemany("""
                INSERT INTO inventory_boards (
                    id_name, ne_type, ne_fdn, ne_name, wifi_device_info, authenticator_name,
                    bios_ver, bios_ver_ex, board_name, board_type, rxu_power_cable_length,
                    clei_code, creditable, creditable_changed_time, date_of_last_service,
                    date_of_manufacture, ext_info, subrack_no, inventory_unit_id,
                    inventory_unit_type, rev_issue_number, pn_bom_code, lan_ver, logic_ver,
                    mbus_ver, manufacturer_data, model, module_no, port_no, port_type,
                    cabinet_no, sn_barcode, slot_no, slot_pos, software_version, subslot_no,
                    unit_position, user_label, vendor_name, vendor_unit_family_type,
                    vendor_unit_type_number, hardware_version
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?);
            """, rows_to_insert)
            conn.commit()
            total += len(rows_to_insert)
            rows_to_insert.clear()
            if total % 50000 == 0:
                print(f"  Processed {total} inventory rows ({time.time()-t0:.1f}s)...")

    if rows_to_insert:
        cursor.executemany("""
            INSERT INTO inventory_boards (
                id_name, ne_type, ne_fdn, ne_name, wifi_device_info, authenticator_name,
                bios_ver, bios_ver_ex, board_name, board_type, rxu_power_cable_length,
                clei_code, creditable, creditable_changed_time, date_of_last_service,
                date_of_manufacture, ext_info, subrack_no, inventory_unit_id,
                inventory_unit_type, rev_issue_number, pn_bom_code, lan_ver, logic_ver,
                mbus_ver, manufacturer_data, model, module_no, port_no, port_type,
                cabinet_no, sn_barcode, slot_no, slot_pos, software_version, subslot_no,
                unit_position, user_label, vendor_name, vendor_unit_family_type,
                vendor_unit_type_number, hardware_version
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?);
        """, rows_to_insert)
        conn.commit()
        total += len(rows_to_insert)

    print(f"Imported {total} rows from {sheet_name} in {time.time()-t0:.2f}s")
    return total

def run_full_etl(excel_path=None):
    """Execute complete ETL process."""
    if not excel_path:
        excel_path = EXCEL_SOURCE_PATH
    
    if not os.path.exists(excel_path):
        raise FileNotFoundError(f"Excel file not found at: {excel_path}")
    
    init_db()
    conn = get_db_connection()
    
    start_time = time.time()
    print(f"Starting ETL from: {excel_path}")
    print("Loading Excel in read-only streaming mode...")
    wb = openpyxl.load_workbook(excel_path, read_only=True, data_only=True)
    
    summary = {}
    summary["gsm"] = import_gsm(wb, conn)
    summary["umts"] = import_umts(wb, conn)
    summary["lte"] = import_lte(wb, conn)
    summary["nr"] = import_nr(wb, conn)
    summary["rru_catalog"] = import_rru_catalog(wb, conn)
    summary["rru_items"] = import_rru_items(wb, conn)
    summary["inventory"] = import_inventory(wb, conn)
    
    wb.close()
    
    total_time = time.time() - start_time
    cursor = conn.cursor()
    cursor.execute("""
        INSERT INTO import_metadata (source_file, completed_at, status, records_summary)
        VALUES (?, CURRENT_TIMESTAMP, 'COMPLETED', ?);
    """, (os.path.basename(excel_path), str(summary)))
    conn.commit()
    conn.close()
    
    print(f"\nETL completed successfully in {total_time:.2f} seconds!")
    print(f"Summary: {summary}")
    return summary

if __name__ == "__main__":
    run_full_etl()
